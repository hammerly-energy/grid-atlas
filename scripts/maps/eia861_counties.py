"""List the RTOs serving each contiguous-US county from EIA-861, for the ISO footprints on the locator map.

EIA-861 lists each utility's counties (Service_Territory) and, per state, its customers and balancing authority
(Sales_Ult_Cust, Delivery_Companies; the "Operating in these RTOs" flags in Utility_Data also count market
participation such as the western EIM, so they are not used). Each utility's
customers are spread evenly over its counties in that state; a county is in an RTO's footprint when that RTO's
utilities hold at least MIN_SHARE of its estimated customers. County-level approximation: a border county can be in
two footprints, and a small utility at the edge of a county doesn't pull the county in.

Also writes each county's NERC regional entity (county_re.csv): the "NERC Region" EIA-861 lists for each
utility in each state, spread over counties the same way; a county goes to the regional entity whose utilities hold
the most of its estimated customers (regional entities don't overlap).

    python3 -I scripts/maps/eia861_counties.py scripts/maps/raw/f8612025   # -> derived/county_rto.csv, derived/county_re.csv
"""
import collections, csv, json, re, sys, unicodedata
from pathlib import Path
import openpyxl

HERE = Path(__file__).resolve().parent
MIN_SHARE = 0.3   # an RTO's utilities must hold at least this share of a county's estimated customers
RTO_BA = {'CISO': 'caiso', 'ERCO': 'ercot', 'PJM': 'pjm', 'MISO': 'miso', 'SWPP': 'spp', 'NYIS': 'nyiso', 'ISNE': 'isone'}   # EIA BA code -> area id
# EIA "NERC Region" (self-reported) -> area id. Legacy codes map to today's regional entity: FRCC joined SERC in 2019;
# SPP RE dissolved in 2018 and most of its registered entities moved to MRO.
NERC_RE = {'WECC': 're_wecc', 'RFC': 're_rf', 'RF': 're_rf', 'SERC': 're_serc', 'FRCC': 're_serc', 'MRO': 're_mro', 'SPP': 're_mro',
           'NPCC': 're_npcc', 'TRE': 're_tre', 'ERCOT': 're_tre'}
FLAG_RE = ['re_tre', 're_serc', 're_mro', 're_npcc', 're_rf', 're_serc', 're_mro', 're_wecc']   # "Also operating in" columns TRE..WECC
FIPS_STATE = {'01': 'AL', '04': 'AZ', '05': 'AR', '06': 'CA', '08': 'CO', '09': 'CT', '10': 'DE', '11': 'DC', '12': 'FL', '13': 'GA',
    '16': 'ID', '17': 'IL', '18': 'IN', '19': 'IA', '20': 'KS', '21': 'KY', '22': 'LA', '23': 'ME', '24': 'MD', '25': 'MA', '26': 'MI',
    '27': 'MN', '28': 'MS', '29': 'MO', '30': 'MT', '31': 'NE', '32': 'NV', '33': 'NH', '34': 'NJ', '35': 'NM', '36': 'NY', '37': 'NC',
    '38': 'ND', '39': 'OH', '40': 'OK', '41': 'OR', '42': 'PA', '44': 'RI', '45': 'SC', '46': 'SD', '47': 'TN', '48': 'TX', '49': 'UT',
    '50': 'VT', '51': 'VA', '53': 'WA', '54': 'WV', '55': 'WI', '56': 'WY'}


def norm(name):
    s = unicodedata.normalize('NFKD', str(name)).encode('ascii', 'ignore').decode().lower()
    s = re.sub(r'\b(saint)\b', 'st', s).replace('.', '').replace("'", '')
    s = re.sub(r'\b(county|parish|borough|city and borough|census area|municipality)\b', '', s)
    return re.sub(r'[^a-z]', '', s)


def main(raw):
    raw = Path(raw)
    topo = json.loads((HERE / 'node_modules/us-atlas/counties-10m.json').read_text())
    fips = {}                                    # (state, normalised name) -> 5-digit FIPS
    for g in topo['objects']['counties']['geometries']:
        st = FIPS_STATE.get(g['id'][:2])
        if st:
            n = norm(g['properties']['name'])
            if st in ('VA', 'MD', 'MO', 'NV') and int(g['id'][2:]) >= 510:    # independent cities share names with counties
                fips[(st, n + 'city')] = g['id']; fips.setdefault((st, n), g['id'])
            else:
                fips[(st, n)] = g['id']

    cust = collections.defaultdict(collections.Counter)    # (utility, state) -> BA code -> customers
    for f, sheet in (('Sales_Ult_Cust_2025.xlsx', 'States'), ('Delivery_Companies_2025.xlsx', 'Delivery Companies')):
        for r in openpyxl.load_workbook(raw / f, read_only=True)[sheet].iter_rows(min_row=4, values_only=True):
            if r[4] in ('Bundled', 'Delivery') and isinstance(r[23], (int, float)):
                cust[(r[1], r[6])][r[8]] += r[23]

    region = {}                                   # (utility, state) -> NERC regional entity
    for r in openpyxl.load_workbook(raw / 'Utility_Data_2025.xlsx', read_only=True)['States'].iter_rows(min_row=3, values_only=True):
        flags = {FLAG_RE[i] for i, v in enumerate(r[6:14]) if v == 'Y'}
        if r[5] in NERC_RE: region[(r[1], r[3])] = NERC_RE[r[5]]
        elif len(flags) == 1: region[(r[1], r[3])] = flags.pop()     # blank region, one region flagged
    by_utility = collections.defaultdict(set)     # multi-state utilities are often listed under one state only
    for (u, _), re_id in region.items(): by_utility[u].add(re_id)

    served = collections.defaultdict(list)        # (utility, state) -> FIPS list
    missing = set()
    for r in openpyxl.load_workbook(raw / 'Service_Territory_2025.xlsx', read_only=True)['Counties_States'].iter_rows(min_row=2, values_only=True):
        f = fips.get((r[4], norm(r[5])))
        if f: served[(r[1], r[4])].append(f)
        elif r[4] in FIPS_STATE.values(): missing.add((r[4], r[5]))

    total = collections.Counter()                 # FIPS -> estimated customers
    by_rto = collections.defaultdict(collections.Counter)    # FIPS -> RTO -> estimated customers
    by_re = collections.defaultdict(collections.Counter)     # FIPS -> regional entity -> estimated customers
    for key, fl in served.items():
        for code, c in cust.get(key, {}).items():
            per = c / len(fl)                     # customers spread evenly over the utility's counties in that state
            for f in fl:
                total[f] += per
                if code in RTO_BA: by_rto[f][RTO_BA[code]] += per
                re_id = region.get(key) or (len(by_utility[key[0]]) == 1 and next(iter(by_utility[key[0]])))
                if re_id: by_re[f][re_id] += per
    out = HERE / 'derived/county_rto.csv'
    weight = {}
    with out.open('w', newline='') as fh:
        w = csv.writer(fh); w.writerow(['fips', 'rto', 'share'])
        for f in sorted(by_rto):
            for rto, c in sorted(by_rto[f].items()):
                if total[f] and c / total[f] >= MIN_SHARE:
                    w.writerow([f, rto, round(c / total[f], 2)]); weight[f] = True
    with (HERE / 'derived/county_re.csv').open('w', newline='') as fh:
        w = csv.writer(fh); w.writerow(['fips', 're', 'share'])
        for f in sorted(by_re):
            re_id, c = by_re[f].most_common(1)[0]
            w.writerow([f, re_id, round(c / total[f], 2)])
    print(f'{len(by_re)} counties with a regional entity: {collections.Counter(c.most_common(1)[0][0] for c in by_re.values())}')
    print(f'{len(weight)} counties in an RTO; {len(missing)} county names unmatched: {sorted(missing)[:15]}')


if __name__ == '__main__':
    main(sys.argv[1])
